"""Summarize EPA's 2022 National Lakes Assessment fish-tissue mercury file.

Usage: uv run python -m pipeline.national <xlsx>
Writes web/public/data/national.json. Run once; the output is committed.
Means are simple (unweighted) means of fillet composites, in ppm.
"""
import json
import statistics
import sys
from collections import Counter, defaultdict

from pipeline.build import OUT

MIN_N = 5
# Column positions in the two sheets of the EPA workbook.
HG_SAMPLE_ID, HG_AMOUNT = 4, 9
INFO_SAMPLE_ID, INFO_COMMON_NAME = 7, 12


def _stat(values):
    return {"mean": round(statistics.mean(values), 4), "n": len(values)}


def summarize(hg_rows, info_rows):
    names = defaultdict(Counter)
    for r in info_rows:
        if r[INFO_COMMON_NAME]:
            names[r[INFO_SAMPLE_ID]][r[INFO_COMMON_NAME]] += 1
    every, by_species = [], defaultdict(list)
    for r in hg_rows:
        if r[HG_AMOUNT] is None:
            continue
        ppm = r[HG_AMOUNT] / 1000  # ng/g -> ppm
        every.append(ppm)
        if names[r[HG_SAMPLE_ID]]:
            species = names[r[HG_SAMPLE_ID]].most_common(1)[0][0].upper()
            by_species[species].append(ppm)
    return {
        "source": "EPA 2022 National Lakes Assessment",
        "all": _stat(every),
        "species": {k: _stat(v) for k, v in sorted(by_species.items()) if len(v) >= MIN_N},
    }


def main():
    import openpyxl

    wb = openpyxl.load_workbook(sys.argv[1], read_only=True)
    hg_sheet, info_sheet = wb.worksheets[0], wb.worksheets[1]
    result = summarize(list(hg_sheet.iter_rows(min_row=2, values_only=True)),
                       list(info_sheet.iter_rows(min_row=2, values_only=True)))
    (OUT / "national.json").write_text(json.dumps(result, sort_keys=True, separators=(",", ":")))
    print(json.dumps(result, indent=1))


if __name__ == "__main__":
    main()
